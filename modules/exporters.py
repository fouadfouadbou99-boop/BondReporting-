from io import BytesIO
from docx import Document
from reportlab.pdfgen import canvas
import pandas as pd


def export_excel(
        synthese,
        portefeuille):

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="xlsxwriter"
    ) as writer:

        synthese.to_excel(
            writer,
            sheet_name="Synthese",
            index=False
        )

        portefeuille.to_excel(
            writer,
            sheet_name="Portefeuille",
            index=False
        )

    return output.getvalue()


def export_word(commentaire):

    output = BytesIO()

    doc = Document()

    doc.add_heading(
        "Reporting Obligataire",
        level=1
    )

    doc.add_paragraph(
        commentaire
    )

    doc.save(output)

    return output.getvalue()


def export_pdf(commentaire):

    output = BytesIO()

    pdf = canvas.Canvas(output)

    text = pdf.beginText(
        40,
        800
    )

    for line in commentaire.split("\n"):

        text.textLine(line)

    pdf.drawText(text)

    pdf.save()

    return output.getvalue()
